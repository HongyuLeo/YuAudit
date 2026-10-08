from __future__ import annotations
import csv
import json
import os
import tempfile
from pathlib import Path
from .models import Result

COLUMNS = ["code", "severity", "identifier", "left", "right", "difference", "rule", "suggestion", "status", "review_note", "sources"]


def safe_cell(value):
    text = str(value)
    # Prevent spreadsheet formula injection from imported document text.
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def finding_rows(result: Result):
    for finding in result.findings:
        data = finding.model_dump()
        data["sources"] = "\n".join(f"{s.location()} = {s.raw}" + (f" [corrected; original={s.original_raw}]" if s.corrected else "") for s in finding.sources)
        yield [safe_cell(data[c]) for c in COLUMNS]


def export_report(result: Result, path: str):
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(dir=target.parent, suffix=target.suffix)
    os.close(handle)
    try:
        if target.suffix.lower() == ".json":
            Path(temporary).write_text(result.model_dump_json(indent=2), encoding="utf-8")
        elif target.suffix.lower() == ".csv":
            with open(temporary, "w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow(COLUMNS)
                writer.writerows(finding_rows(result))
        elif target.suffix.lower() == ".xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from .models import FIELDS
            book = Workbook()
            summary = book.active
            summary.title = "Summary"
            summary.append(["YuAudit", result.version])
            for k, v in {"created_at": result.created_at, "conclusion": result.conclusion,
                         "deterministic_checks": result.checks, "findings": len(result.findings),
                         "rules": result.rules.model_dump_json(), "AI draft (unverified)": result.ai_summary or "disabled"}.items():
                summary.append([k, safe_cell(v)])
            exceptions = book.create_sheet("Exceptions")
            exceptions.append(COLUMNS)
            for row in finding_rows(result):
                exceptions.append(row)
            data = book.create_sheet("Structured data")
            data.append(["file", "kind", "partial", *FIELDS])
            evidence = book.create_sheet("Evidence")
            evidence.append(["file", "kind", "sku", "field", "value", "location", "bbox", "original_raw", "corrected"])
            for doc in result.documents:
                for line in doc.lines:
                    data.append([safe_cell(doc.path), doc.kind.value, doc.partial,
                                 *[safe_cell(getattr(line, f) if getattr(line, f) is not None else "") for f in FIELDS]])
                    for field, source in line.evidence.items():
                        evidence.append([safe_cell(doc.path), doc.kind.value, safe_cell(line.sku or ""), field,
                                         safe_cell(source.raw), safe_cell(source.location()), str(source.bbox or ""),
                                         safe_cell(source.original_raw or ""), source.corrected])
            for sheet in book:
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="334155")
                for col in sheet.columns:
                    letter = col[0].column_letter
                    sheet.column_dimensions[letter].width = min(64, max(16, max(len(str(c.value or "")) for c in col[:200]) + 2))
                for row in sheet.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
            book.save(temporary)
        else:
            raise ValueError("Export requires .xlsx / .csv / .json")
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
