from decimal import Decimal
from pathlib import Path
from threading import Event
import csv
import json
import pytest
from yuaudit.models import Document, Kind, Line, Rules, decimal_value, Finding, Source
from yuaudit.parsers import read_document, Cancelled
from yuaudit.normalize import structure, guess_mapping
from yuaudit.engine import reconcile
from yuaudit.demo import create_demo
from yuaudit.reports import export_report, safe_cell
from yuaudit.ai import MappingAdvice, OllamaProvider


def docs(a="10", b="10", price="0.10", currency="USD", partial=False):
    base = dict(po_number="PO1", sku="SKU1", quantity=Decimal(a), unit_price=Decimal("0.10"), currency="USD")
    po = Document(id="p", path="po.csv", kind=Kind.PO, lines=[Line(**base)])
    inv = Line(**{**base, "quantity": Decimal(b), "unit_price": Decimal(price),
                  "line_total": Decimal(b) * Decimal(price), "invoice_number": "INV1", "currency": currency})
    pack = Line(po_number="PO1", invoice_number="INV1", sku="SKU1", quantity=Decimal(b))
    return [po, Document(id="i", path="invoice.csv", kind=Kind.INVOICE, lines=[inv], partial=partial),
            Document(id="k", path="packing.csv", kind=Kind.PACKING, lines=[pack])]


def codes(result):
    return [f.code for f in result.findings]


@pytest.mark.parametrize("value,expected", [("1,234.56", "1234.56"), ("0.10", "0.10"), ("", None), ("-1.2", "-1.2")])
def test_decimal(value, expected):
    assert decimal_value(value) == (Decimal(expected) if expected else None)


@pytest.mark.parametrize("value", ["1,23", "NaN", "Infinity", "1e9", "0.123456789", "1,2.3"])
def test_invalid_decimal(value):
    with pytest.raises(ValueError):
        decimal_value(value)


def test_exact_money():
    result = reconcile(docs(a="3", b="3"))
    assert result.findings == []
    assert result.conclusion == "no_exceptions_in_checked_scope"


def test_quantity_price_and_arithmetic():
    data = docs(b="9", price="0.12")
    data[1].lines[0].line_total = Decimal("9.00")
    result = reconcile(data)
    assert {"quantity", "unit_price", "arithmetic"} <= set(codes(result))
    assert next(f for f in result.findings if f.code == "unit_price").difference == "0.02"


def test_partial_under_not_error_over_is_error():
    result = reconcile(docs(a="100", b="40", partial=True))
    assert "partial_pending" in codes(result)
    assert "quantity" not in codes(result)
    assert "quantity" in codes(reconcile(docs(a="100", b="101", partial=True)))


def test_multi_invoice_batches():
    data = docs(a="100", b="40")
    second = data[1].lines[0].model_copy(update={"quantity": Decimal(60), "line_total": Decimal(6), "invoice_number": "INV2"})
    data[1].lines.append(second)
    data[2].lines.append(Line(po_number="PO1", invoice_number="INV2", sku="SKU1", quantity=Decimal(60)))
    assert reconcile(data).findings == []


@pytest.mark.parametrize("currency", ["EUR", None])
def test_currency_blocks_price(currency):
    result = reconcile(docs(price="999", currency=currency))
    assert "currency" in codes(result)
    assert "unit_price" not in codes(result)


def test_missing_is_not_pass():
    data = docs()
    data[1].lines[0].quantity = None
    assert {"missing_field", "unverified"} <= set(codes(reconcile(data)))
    assert "unverified" in codes(reconcile([]))


def test_duplicate_and_invoice_reference():
    data = docs()
    data[1].lines.append(data[1].lines[0].model_copy())
    assert {"duplicate_line", "unverified"} <= set(codes(reconcile(data)))
    other = data[1].model_copy(deep=True)
    other.id = "other"
    assert "duplicate_invoice" in codes(reconcile(data + [other]))
    data = docs()
    data[1].lines[0].po_number = "OTHER"
    assert "document_relation" in codes(reconcile(data))


def test_fuzzy_never_automatic():
    data = docs()
    data[1].lines[0].sku = "SKU-I"
    result = reconcile(data, Rules(fuzzy_threshold=0.5))
    assert "sku_mismatch" in codes(result)
    assert result.conclusion == "review_required"


def test_normalization_collision():
    data = docs()
    data[0].lines.append(data[0].lines[0].model_copy(update={"sku": "SKU-1"}))
    assert "ambiguous_match" in codes(reconcile(data))


def test_po_reference_collision_reproduction():
    data = docs(a="10", b="30", price="1")
    for doc in data:
        doc.lines[0].po_number = "PO-1"
        doc.lines[0].unit_price = Decimal(1)
    data[0].lines.append(data[0].lines[0].model_copy(deep=True, update={"po_number": "PO1", "quantity": Decimal(20)}))
    for index, line in enumerate(data[0].lines, 2):
        line.evidence["po_number"] = Source(file="po.csv", row=index, column="PO Number", raw=line.po_number)
    result = reconcile(data)
    collision = next(f for f in result.findings if f.code == "ambiguous_match")
    assert result.conclusion == "review_required"
    assert collision.left == "PO-1 | PO1"
    assert {s.row for s in collision.sources} == {2, 3}
    assert not {"quantity", "unit_price"} & set(codes(result))
    assert result.checks == 1  # only invoice row arithmetic; no aggregate comparisons


@pytest.mark.parametrize("field,first,second", [
    ("po_number", "PO-1", "PO1"),
    ("invoice_number", "INV-1", "INV1"),
    ("sku", "SKU-1", "SKU1"),
])
@pytest.mark.parametrize("kind_index", [0, 1, 2])
def test_identifier_collisions_across_documents_block_aggregation(field, first, second, kind_index):
    data = docs()
    data[kind_index].lines[0].__setattr__(field, first)
    extra = data[kind_index].model_copy(deep=True)
    extra.id += "-second"
    extra.path = "second.csv"
    extra.lines[0].__setattr__(field, second)
    result = reconcile(data + [extra])
    assert result.conclusion == "review_required"
    assert any(f.code == "ambiguous_match" and field in f.rule for f in result.findings)
    assert not {"quantity", "unit_price"} & set(codes(result))
    if field == "invoice_number":
        assert "duplicate_invoice" not in codes(result)  # ambiguity is not proof of repeated invoicing


@pytest.mark.parametrize("field", ["po_number", "invoice_number", "sku"])
@pytest.mark.parametrize("normalization", [True, False])
def test_empty_normalized_identifiers_never_match(field, normalization):
    data = docs()
    for index, doc in enumerate(data):
        setattr(doc.lines[0], field, "-" * (index + 1))
    result = reconcile(data, Rules(normalize_identifiers=normalization))
    assert result.conclusion == "review_required"
    assert sum(f.code == "ambiguous_match" for f in result.findings) == 3
    assert result.checks == 1


def test_unambiguous_identifier_normalization_all_fields():
    data = docs()
    data[0].lines[0].po_number = "PO-1"
    data[0].lines[0].sku = "SKU-1"
    data[1].lines[0].invoice_number = "INV-1"
    data[2].lines[0].invoice_number = "inv1"
    assert reconcile(data).findings == []


def test_collision_does_not_block_unrelated_groups():
    data = docs(b="9")
    for sku in ("OTHER-1", "OTHER1"):
        data[0].lines.append(data[0].lines[0].model_copy(update={"sku": sku}))
    result = reconcile(data)
    assert {"ambiguous_match", "quantity"} <= set(codes(result))
    assert next(f for f in result.findings if f.code == "quantity").identifier == "SKU1"


def test_collision_blocks_declared_total():
    data = docs()
    data[1].lines[0].total_amount = Decimal(2)
    data[1].lines.append(data[1].lines[0].model_copy(update={"sku": "SKU-1"}))
    result = reconcile(data)
    assert any(f.code == "unverified" and "total aggregation" in f.rule for f in result.findings)
    assert "total_amount" not in codes(result)


def test_total():
    data = docs()
    data[1].lines[0].total_amount = Decimal("2.00")
    assert "total_amount" in codes(reconcile(data))


def test_cancel():
    event = Event()
    event.set()
    with pytest.raises(Cancelled):
        reconcile(docs(), cancel=event)


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "gbk"])
def test_csv_chinese_path(tmp_path, encoding):
    path = tmp_path / "中文 用户 空格" / "采购订单.csv"
    path.parent.mkdir()
    path.write_text("订单号,货号,数量,单价,币种\nPO1,A1,3,0.10,USD\n", encoding=encoding)
    before = path.read_bytes()
    doc = structure(read_document(str(path)))
    assert doc.lines[0].quantity == Decimal(3)
    assert doc.lines[0].evidence["sku"].row == 2
    assert path.read_bytes() == before


def test_xlsx(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "PO.xlsx"
    book = Workbook()
    book.active.title = "采购"
    book.active.append(["PO Number", "SKU", "Quantity"])
    book.active.append(["PO1", "A", 3])
    sheet = book.create_sheet("Second")
    sheet.append(["PO Number", "SKU", "Quantity"])
    sheet.append(["PO1", "B", 2])
    book.save(path)
    doc = structure(read_document(str(path)))
    assert len(doc.lines) == 2
    assert doc.lines[0].evidence["sku"].sheet == "采购"


def test_pdf_and_scanned(tmp_path):
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.pdfgen.canvas import Canvas
    path = tmp_path / "PO.pdf"
    table = Table([["PO Number", "SKU", "Quantity"], ["PO1", "A1", "3"]])
    table.setStyle(TableStyle([("GRID", (0,0), (-1,-1), 1, colors.black)]))
    SimpleDocTemplate(str(path)).build([table])
    doc = structure(read_document(str(path)))
    assert doc.lines[0].quantity == Decimal(3)
    assert doc.lines[0].evidence["sku"].page == 1
    assert doc.lines[0].evidence["sku"].bbox
    scan = tmp_path / "scan.pdf"
    canvas = Canvas(str(scan))
    canvas.rect(10, 10, 100, 100)
    canvas.showPage()
    canvas.save()
    with pytest.raises(ValueError, match="OCR"):
        read_document(str(scan))


def test_mapping_invalid(tmp_path):
    paths = create_demo(tmp_path)
    doc = read_document(paths[0])
    doc.mapping["SKU"] = "quantity"
    with pytest.raises(ValueError, match="Multiple"):
        structure(doc)


def test_corrections_keep_original(tmp_path):
    doc = read_document(create_demo(tmp_path)[0])
    doc.rows[0].values["Quantity"] = "11"
    line = structure(doc).lines[0]
    assert line.evidence["quantity"].corrected
    assert line.evidence["quantity"].original_raw == "10"
    assert line.evidence["quantity"].raw == "11"


def test_tolerance_and_strategy():
    data = docs(b="9.999", price="0.109")
    assert "quantity" in codes(reconcile(data))
    assert "quantity" not in codes(reconcile(data, Rules(quantity_tolerance="0.01")))
    data = docs()
    data[1].lines[0].sku = "sku-1"
    assert "extra_invoice_item" not in codes(reconcile(data))
    assert "extra_invoice_item" in codes(reconcile(data, Rules(normalize_identifiers=False, fuzzy_enabled=False)))


def test_bad_formats_and_cancel_parser(tmp_path):
    path = tmp_path / "broken.csv"
    path.write_text("SKU,SKU\nA,B", encoding="utf-8")
    with pytest.raises(ValueError, match="Duplicate"):
        read_document(str(path))
    event = Event()
    event.set()
    with pytest.raises(Cancelled):
        read_document(create_demo(tmp_path)[0], cancel=event)


def test_cross_po_never_matches():
    data = docs()
    data[1].lines[0].po_number = "PO2"
    result = reconcile(data)
    assert "document_relation" in codes(result)
    assert "unit_price" not in codes(result)


def test_duplicate_po_across_files():
    data = docs()
    duplicate = data[0].model_copy(deep=True)
    duplicate.id = "p2"
    result = reconcile(data + [duplicate])
    assert {"duplicate_line", "unverified"} <= set(codes(result))
    assert "quantity" not in codes(result)


def test_demo_expected_and_reports(tmp_path):
    data = [structure(read_document(p)) for p in create_demo(tmp_path / "虚构 数据")]
    result = reconcile(data)
    expected = {"quantity", "unit_price", "missing_invoice_item", "extra_invoice_item", "duplicate_line", "sku_mismatch", "unverified"}
    assert set(codes(result)) == expected
    assert not any(f.identifier in ("GOOD-01", "BATCH-07") for f in result.findings)
    assert len(result.findings) == 8
    for suffix in ("xlsx", "csv", "json"):
        path = tmp_path / f"报告.{suffix}"
        export_report(result, str(path))
        assert path.stat().st_size > 100
    from openpyxl import load_workbook
    book = load_workbook(tmp_path / "报告.xlsx")
    assert book.sheetnames == ["Summary", "Exceptions", "Structured data", "Evidence"]
    assert book["Exceptions"].max_row == 9
    payload = json.loads((tmp_path / "报告.json").read_text(encoding="utf-8"))
    assert payload["findings"][0]["sources"]


def test_injection():
    assert safe_cell("=HYPERLINK(\"x\")").startswith("'")
    assert safe_cell(" @cmd").startswith("'")


def test_ai_schema_and_local_only():
    with pytest.raises(ValueError):
        MappingAdvice(mapping={"x": "invented"}, explanation="x")
    with pytest.raises(ValueError):
        OllamaProvider("https://example.com")
    OllamaProvider()


def test_storage(tmp_path, monkeypatch):
    monkeypatch.setenv("YUAUDIT_DATA_DIR", str(tmp_path / "本地 缓存"))
    from yuaudit.storage import save_session, load_session
    result = reconcile(docs())
    save_session(result)
    assert load_session() == result
